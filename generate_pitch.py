from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Define simple styling helper
    def set_title(slide, text):
        title = slide.shapes.title
        title.text = text
        for p in title.text_frame.paragraphs:
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 51, 102)

    # 1. Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Smart QR Attendance & ERP System"
    subtitle.text = "Secure, Scalable, and Geofenced Attendance Management\n\nFinal Year Project Presentation"
    
    # 2. Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "The Problem")
    content = slide.placeholders[1].text_frame
    content.text = "Current attendance systems face major challenges:"
    p = content.add_paragraph()
    p.text = "Time Consuming: Calling roll numbers wastes 10-15 minutes of lecture time."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Proxy Attendance: Manual signatures or basic apps are easily manipulated by students."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Data Silos: Attendance data is rarely synchronized with the college's main ERP in real-time."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Poor Analytics: Hard to identify students with critically low attendance instantly."
    p.level = 1

    # 3. Core Features
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Core Features & Solutions")
    content = slide.placeholders[1].text_frame
    p = content.add_paragraph()
    p.text = "Dynamic QR Codes: QR codes refresh every 10 seconds to prevent screenshots."
    p = content.add_paragraph()
    p.text = "Geofencing Validation: Students must be physically within the classroom's GPS radius (e.g. 50 meters) to scan."
    p = content.add_paragraph()
    p.text = "Device Fingerprinting: Prevents a student from marking attendance for multiple accounts on a single phone."
    p = content.add_paragraph()
    p.text = "Deep RBAC Hierarchy: Strict separation of concerns (Principal > HOD > Class Teacher > Faculty > Student)."
    p = content.add_paragraph()
    p.text = "Live Analytics Dashboard: Real-time graphs for HODs and Principals to monitor attendance trends."

    # 4. Worktree / Hierarchy
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "System Hierarchy (Worktree)")
    content = slide.placeholders[1].text_frame
    content.text = "Data and visibility flow securely from top to bottom:"
    
    p = content.add_paragraph()
    p.text = "1. Super Admin (Manages multiple colleges & system health)"
    p.level = 1
    p = content.add_paragraph()
    p.text = "2. Principal (Global view of their specific college)"
    p.level = 2
    p = content.add_paragraph()
    p.text = "3. HOD (Head of Department)"
    p.level = 3
    p = content.add_paragraph()
    p.text = "Views all Faculty & Class Teachers within their branch"
    p.level = 4
    p = content.add_paragraph()
    p.text = "4. Class Teacher / Faculty"
    p.level = 4
    p = content.add_paragraph()
    p.text = "Manages attendance, timetable, and analytics for their assigned classes"
    p.level = 5
    p = content.add_paragraph()
    p.text = "5. Student (Views own attendance and scans QR)"
    p.level = 6

    # 5. Flowchart of Attendance
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "QR Scanning Flowchart")
    content = slide.placeholders[1].text_frame
    content.text = "Step-by-step process of secure attendance marking:"
    
    p = content.add_paragraph()
    p.text = "[START] Faculty generates session on Smart Board"
    p.level = 1
    p = content.add_paragraph()
    p.text = "➔ QR Code appears and rotates every 10s"
    p.level = 2
    p = content.add_paragraph()
    p.text = "[SCAN] Student scans QR via Mobile App"
    p.level = 1
    p = content.add_paragraph()
    p.text = "➔ Check 1: Is the token valid and not expired?"
    p.level = 2
    p = content.add_paragraph()
    p.text = "➔ Check 2: Is student's GPS within classroom radius?"
    p.level = 2
    p = content.add_paragraph()
    p.text = "➔ Check 3: Is device fingerprint unique?"
    p.level = 2
    p = content.add_paragraph()
    p.text = "[SUCCESS] Attendance marked in database instantly."
    p.level = 1

    # 6. Third-Party ERP Integration
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Third-Party ERP Integration Strategy")
    content = slide.placeholders[1].text_frame
    content.text = "How we integrate with a College's existing legacy ERP:"
    
    p = content.add_paragraph()
    p.text = "RESTful API Endpoints: We provide secure API keys for the college to fetch daily attendance logs."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Webhooks: The system can instantly POST attendance data to the college's ERP server the moment a session ends."
    p.level = 1
    p = content.add_paragraph()
    p.text = "CSV/Excel Exports: Standardized automated reports generated daily for manual ERP imports."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Data Preservation: Our system acts as the 'Attendance Microservice', ensuring the legacy ERP remains the ultimate source of truth."
    p.level = 1

    # 7. Business Model
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Business Model (B2B SaaS)")
    content = slide.placeholders[1].text_frame
    content.text = "A sustainable path to commercialization:"
    
    p = content.add_paragraph()
    p.text = "Tier 1: Free/Basic (Open for small institutes, max 100 students) - Serves as product marketing."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Tier 2: Pro Tier ($X per student/year) - Unlocks Geofencing, Advanced Analytics, and ERP Webhooks."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Tier 3: Enterprise White-label - We deploy a custom-branded version of the app directly on the college's own AWS/Azure servers for strict data compliance."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Hardware Upsell: Selling dedicated RFID/QR kiosk tablets for classroom doors."
    p.level = 1

    # 8. Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Conclusion")
    content = slide.placeholders[1].text_frame
    content.text = "Summary:"
    p = content.add_paragraph()
    p.text = "A robust, proxy-proof solution solving a real-world problem."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Modern Tech Stack ensures speed, reliability, and security."
    p.level = 1
    p = content.add_paragraph()
    p.text = "Highly scalable with a clear path to API integration and commercial viability."
    p.level = 1

    prs.save('Attendance_System_Pitch.pptx')

if __name__ == '__main__':
    create_presentation()
